#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
超声医学报告批量质控与错误分析脚本
====================================
功能：
    1. 遍历 init_report/ 下所有 .txt 超声报告；
    2. 调用 DeepSeek API（deepseek-chat，成本最低的模型）对每份报告做多维度质控审核；
    3. 多线程并发 + 自动重试 + tqdm 进度条；
    4. 结果导出 Excel（质控概览 + 错误明细表 两个 Sheet，概览含错误类型柱状图）。

用法：
    python ultrasound_qc.py                 # 处理全部报告
    python ultrasound_qc.py --limit 20      # 先试跑 20 份，验证效果与成本
    python ultrasound_qc.py --workers 15    # 自定义并发数
    python ultrasound_qc.py --out result.xlsx

依赖安装：
    pip install openai pandas openpyxl tqdm python-dotenv
"""

import argparse
import json
import os
import random
import re
import sys
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import pandas as pd
from openai import (
    APIConnectionError,
    APITimeoutError,
    InternalServerError,
    OpenAI,
    RateLimitError,
)
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

# ============================ 配置区 ============================
# 密钥通过环境变量或 .env 文件提供（DEEPSEEK_API_KEY=sk-xxx），严禁硬编码到代码仓库
API_KEY = os.getenv("DEEPSEEK_API_KEY", "")
BASE_URL = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com")

# 模型选择：deepseek-chat 为性价比最高档（约为 deepseek-reasoner 价格的几分之一），
# 本任务为规则明确的报告审核，无需深度推理，deepseek-chat 足够且最省 token
MODEL_NAME = "deepseek-chat"

REPORT_DIR = Path("init_report")        # 报告文件夹
OUTPUT_XLSX = "ultrasound_qc_report.xlsx"  # 输出 Excel
DEFAULT_WORKERS = 10                    # 默认并发线程数（10-20）
MAX_RETRIES = 3                         # API 最大重试次数
TEMPERATURE = 0.1                       # 低温度保证审核稳定性
MAX_TOKENS = 2000                       # 单次响应 token 上限
REQUEST_TIMEOUT = 90                    # 单次请求超时（秒）
FLAG_AGE_FORMAT = False                 # 是否将年龄前导零格式（如"040Y"）计为错误（False=忽略，减少噪音）

# 可选：读取 .env 文件（未安装 python-dotenv 则跳过）
try:
    from dotenv import load_dotenv
    load_dotenv()
    API_KEY = os.getenv("DEEPSEEK_API_KEY", API_KEY)
    BASE_URL = os.getenv("DEEPSEEK_BASE_URL", BASE_URL)
except ImportError:
    pass

# 标准错误分类（用于汇总统计，模型输出会归一化到这些分类）
# 分类依据：国家卫健委《超声诊断专业医疗质量控制指标（2022年版）》US-RE-01 超声报告书写合格率
# 对"不合格报告"的界定，以及 US-RE-BR-01（BI-RADS 分类率）、C-TIRADS（2020 中国甲状腺指南）
VALID_CATEGORIES = {
    "描述与结论不一致",   # US-RE-01 第3条：报告单中的描述与结论不一致
    "遗漏检查部位",       # US-RE-01 第2条：未包含申请单开具项目检查的
    "方位侧别错误",       # US-RE-01 第4条：器官/部位/病变的方位（左右、上下、前后）错误
    "性别与检查冲突",     # US-RE-01 第4条：性别与检查部位/器官不符
    "性别年龄冲突",       # US-RE-01 第4条：年龄等基本信息与检查内容不符
    "脏器缺如报告正常",   # US-RE-01 第4条：所查脏器缺如但报告为正常
    "模板文字未删除",     # US-RE-01 第4条：未删除与超声报告有歧义的模板文字
    "数据或单位错误",     # US-RE-01 第4条：单位、数据错误
    "分级不规范",         # US-RE-BR-01：BI-RADS 用于乳腺、甲状腺用 C-TIRADS，且分级须与描述相符
    "错别字或字段缺失",   # 报告完整性（国家超声质控中心多中心研究"报告完整率"维度）
    "语病或表述不规范",   # 基础书写质量：病句、不通顺、口语化、术语不当、标点错误
    "其他",
}

# ============================ 提示词 ============================
SYSTEM_PROMPT = """你是资深超声医学科报告质控专家。请依据国家卫生健康委《超声诊断专业医疗质量控制指标（2022年版）》中"超声报告书写合格率（US-RE-01）"对不合格报告的界定，结合中华医学会超声医学分会相关规范（乳腺BI-RADS、甲状腺C-TIRADS分级体系），审核超声报告质量。

审核维度与标准分类（error_categories 与 error_details[].type 必须严格使用以下分类词）：
1. "描述与结论不一致"：超声所见(finding)与诊断结论(conclusion)不对应——所见描述的病变结论遗漏、结论中的病变在所见中无依据、回声性质/大小/数量/部位指代等描述前后矛盾或存在歧义。（US-RE-01第3条）
2. "遗漏检查部位"：检查项目(examItem)明确列出的器官/部位，在超声所见(finding)和诊断结论(conclusion)中均完全无相应描述（如检查项目为"肝胆胰脾"而所见与结论均无胰腺、脾脏的描述）。注意：仅对检查项目中明确写出的具体器官逐一核对；检查项目为泛称（如"腹部""泌尿系""产前常规""彩超"）时不强拆器官；检查项目为空则跳过此项。（US-RE-01第2条：未包含申请单开具项目检查的）
3. "方位侧别错误"：器官、部位、病变的方位（左右、上下、前后）前后矛盾，如所见描述"左侧"而结论写"右侧"。（US-RE-01第4条）
4. "性别与检查冲突"：报告性别与检查所含器官/部位明确矛盾——女性报告出现前列腺、精囊、睾丸等男性器官的描述或结论，男性报告出现子宫、附件、卵巢、妊娠、哺乳等女性器官/状态的描述或结论。（US-RE-01第4条：报告单患者性别与实际不符）
5. "性别年龄冲突"：年龄与检查内容在医学上明确矛盾（如报告年龄与胎儿状态、绝经后表述明显不符）。注意：儿童/青少年进行乳腺、生殖器、腹部等超声检查临床常见（发育评估、隐睾排查、腹痛检查等），不要仅凭年龄小判定错误。
6. "脏器缺如报告正常"：患者脏器已切除或先天缺如（如报告提及"胆囊切除术后""左肾切除术后"），但所见/结论仍描述该脏器大小形态正常。（US-RE-01第4条）
7. "模板文字未删除"：套用报告模板后未删除的冗余或自相矛盾文字，如同一部位先描述"探及肿大淋巴结"紧接一行"未见肿大淋巴结"、正常器官模板段落与实际病变描述并存造成歧义。（US-RE-01第4条）
8. "数据或单位错误"：测量单位错误（如将cm写成mm且数值不符）、同一病灶前后数据不一致、孕周与生物测量明显不符（差异超过2周）。（US-RE-01第4条）
9. "分级不规范"：仅在分级与声像图特征存在实质矛盾或明显违反通行标准时报此分类——如描述边界不清、形态不规则、微钙化等可疑恶性征象而乳腺评BI-RADS 2/3类；结节内见微钙化而甲状腺评3类以下；仅笼统写"4类"未细分4A/4B/4C；胎盘成熟度Grannum分级使用不存在的"II级+"；肺动脉收缩压实测值与结论的轻/中/重度分级明显矛盾。以下情况一律不报告：甲状腺使用"TI-RADS"缩写本身（TI-RADS即甲状腺影像报告和数据系统，是甲状腺专用体系，C-TIRADS与ACR TI-RADS均为其版本）；正常甲状腺写TI-RADS 1类（C-TIRADS 1类即阴性/正常）；前列腺超声报告未使用PI-RADS（PI-RADS仅用于前列腺mpMRI）；除乳腺外其他器官超声报告未做RADS类分级；描述特征缺失而称"无法验证分级"（这不是错误证据）；"轻中度"反流等临床惯用表述；室壁运动"评分/等级"用词差异。
10. "错别字或字段缺失"：明显错别字、漏字、词语错误，或缺少性别、年龄、超声所见、诊断结论等关键字段（报告完整性）。
11. "语病或表述不规范"：基础书写文字错误——病句或句子成分残缺导致语义不通、口语化/随意表述（如"没事""长了个东西"）、医学术语使用明显不当或自造词语、标点符号错误严重影响理解、同一句话重复啰嗦。（报告书写文字质量；注意：行业普遍接受的习惯简称不算错误）
12. "其他"：以上未涵盖的明确质量问题。

判定注意事项：
a) 审慎判断：只报告有明确依据的问题，不臆测、不过度报告；每个具体问题对应一条error_details记录，同一问题不重复报告。
b) 临界值/正常高值（如颈动脉内-中膜0.8-1.0mm、少量瓣膜反流）且结论为"未见明显异常"或"定期复查"，属临床可接受判断，不报告。
c) 孕周/胎儿生长综合评估：多项生物测量对应孕周小幅差异（2周以内）属正常生理波动，结论采用综合孕周或平均值是规范做法，不报告。
d) 明确病理性异常被结论遗漏必须报告（如E/A比值<1提示左室舒张功能减低而结论遗漏、所见描述占位而结论未提及）。
f) original_text 必须摘自原文且逐字一致；引用跨行内容时按原文分行引用，不要改写、合并或缩写。

输出要求：
1. 只输出一个JSON对象，禁止输出任何解释文字、markdown代码块或其他多余内容。
2. JSON结构如下：
{
  "has_error": true或false,
  "error_categories": ["错误分类"],
  "error_details": [
    {
      "type": "错误分类",
      "original_text": "原文中存在问题的句子（必须摘自原文）",
      "reason": "错误原因分析",
      "suggestion": "修改建议"
    }
  ]
}
3. 报告无问题时：has_error为false，error_categories与error_details均为空数组[]。"""

if FLAG_AGE_FORMAT:
    SYSTEM_PROMPT += ('\ng) 年龄字段的前导零格式（如"040Y"）应报告，归入"错别字或字段缺失"。')
else:
    SYSTEM_PROMPT += ('\ne) 年龄前导零格式（如"040Y"）属于录入格式不规范，'
                      "默认不作为错误报告，除非造成医学歧义。")

# ============================ 报告解析 ============================
FIELD_KEYS = ("examItem", "checkTime", "gender", "age", "finding", "conclusion")


def read_text_anyencoding(path: Path) -> str:
    """依次尝试常见编码读取文本，避免 Windows 下 GBK/UTF-8 混杂导致读取失败。"""
    for enc in ("utf-8", "gbk", "utf-16"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def parse_report(text: str) -> dict:
    """解析 key: value 格式报告，value 支持多行（直到出现下一个已知字段）。"""
    fields = {}
    current = None
    for raw_line in text.splitlines():
        m = re.match(r"^\s*([A-Za-z_]\w*)\s*:\s*(.*)$", raw_line)
        if m and m.group(1) in FIELD_KEYS:
            current = m.group(1)
            fields[current] = m.group(2).strip()
        elif current:
            line = raw_line.strip()
            if line:
                fields[current] = (fields.get(current, "") + "\n" + line).strip()
    return fields


def build_user_prompt(fields: dict, raw_text: str) -> str:
    """将解析后的字段拼装为用户提示词；解析失败则退回原始全文。"""
    if not fields.get("finding") and not fields.get("conclusion"):
        return f"请审核以下超声报告并按要求输出JSON：\n{raw_text}"
    lines = [
        "请审核以下超声报告并按要求输出JSON。",
        f"【检查项目】{fields.get('examItem') or '（缺失）'}",
        f"【检查时间】{fields.get('checkTime') or '（缺失）'}",
        f"【性别】{fields.get('gender') or '（缺失）'}",
        f"【年龄】{fields.get('age') or '（缺失）'}",
        f"【超声所见】{fields.get('finding') or '（缺失）'}",
        f"【诊断结论】{fields.get('conclusion') or '（缺失）'}",
    ]
    return "\n".join(lines)


# ============================ API 调用 ============================
RETRYABLE_ERRORS = (APIConnectionError, APITimeoutError, RateLimitError, InternalServerError)


def call_with_retry(fn, *args, **kwargs):
    """指数退避重试：网络超时/连接错误/限流/服务端5xx 时自动重试。"""
    last_exc = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            return fn(*args, **kwargs)
        except RETRYABLE_ERRORS as e:
            last_exc = e
            sleep_s = 2 ** attempt + random.uniform(0, 1)
            print(f"\n[重试 {attempt}/{MAX_RETRIES}] {type(e).__name__}，{sleep_s:.1f}s 后重试...", flush=True)
            time.sleep(sleep_s)
    raise last_exc


def extract_json(content: str) -> dict:
    """从模型输出中稳健提取 JSON（容忍 markdown 围栏等杂质）。"""
    content = content.strip()
    content = re.sub(r"^```(?:json)?", "", content).strip()
    content = re.sub(r"```$", "", content).strip()
    start, end = content.find("{"), content.rfind("}")
    if start == -1 or end == -1 or end <= start:
        raise ValueError(f"响应中未找到JSON：{content[:100]!r}")
    body = content[start:end + 1]
    # strict=False 容忍字符串中的换行/制表符等控制字符
    try:
        return json.loads(body, strict=False)
    except json.JSONDecodeError:
        # 输出被 max_tokens 截断：截到最后一个完整对象后闭合数组与根对象
        repaired = _repair_truncated_json(body)
        if repaired is None:
            raise
        return json.loads(repaired, strict=False)


def _repair_truncated_json(body: str):
    """修复被截断的 JSON：丢弃最后半个对象，补齐 ']' 与 '}'。"""
    last = body.rfind("}")
    if last <= 0:
        return None
    t = body[:last + 1]
    if t.count('"') % 2:  # 括号前有未闭合引号（如截断在键名中）
        t = t[:last].rstrip().rstrip(",")
        last = t.rfind("}")
        if last <= 0:
            return None
        t = t[:last + 1]
        if t.count('"') % 2:
            return None
    if t.count("[") > t.count("]"):
        t += "]"
    if t.count("{") > t.count("}"):
        t += "}"
    return t


def normalize_result(data: dict) -> dict:
    """校验并归一化模型输出：分类收敛到标准词、类型纠偏、一致性修补。"""
    raw_flag = data.get("has_error")
    has_error = raw_flag in (True, "true", "True", 1, "1")

    details_raw = data.get("error_details")
    details = []
    if isinstance(details_raw, list):
        for d in details_raw:
            if not isinstance(d, dict):
                continue
            t = str(d.get("type") or "其他").strip()
            if t not in VALID_CATEGORIES:
                t = "其他"
            details.append({
                "type": t,
                "original_text": str(d.get("original_text") or "").strip(),
                "reason": str(d.get("reason") or "").strip(),
                "suggestion": str(d.get("suggestion") or "").strip(),
            })

    cats_raw = data.get("error_categories")
    cats = []
    if isinstance(cats_raw, list):
        cats = [str(c).strip() for c in cats_raw if str(c).strip()]
    cats = [c if c in VALID_CATEGORIES else "其他" for c in cats]

    # 一致性修补
    if details and not has_error:
        has_error = True
    if has_error and not cats:
        cats = sorted({d["type"] for d in details})
    if not has_error:
        details, cats = [], []
    return {"has_error": has_error, "error_categories": cats, "error_details": details}


def qc_one_report(client: OpenAI, path: Path, model: str) -> dict:
    """审核单份报告。任何异常都在此捕获并记录，不影响主流程。"""
    result = {
        "report_id": path.stem,
        "status": "ok",
        "error_msg": "",
        "exam_item": "", "gender": "", "age": "", "check_time": "",
        "has_error": False,
        "error_categories": [],
        "error_details": [],
    }
    try:
        raw = read_text_anyencoding(path)
        fields = parse_report(raw)
        result.update(
            exam_item=fields.get("examItem", ""),
            gender=fields.get("gender", ""),
            age=fields.get("age", ""),
            check_time=fields.get("checkTime", ""),
        )
        user_prompt = build_user_prompt(fields, raw)
        resp = call_with_retry(
            client.chat.completions.create,
            model=model,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": user_prompt},
            ],
            temperature=TEMPERATURE,
            max_tokens=MAX_TOKENS,
            response_format={"type": "json_object"},  # DeepSeek 支持 JSON 输出模式
        )
        content = resp.choices[0].message.content or ""
        result.update(normalize_result(extract_json(content)))
    except Exception as e:  # noqa: BLE001 单份失败不影响整体
        result["status"] = "failed"
        result["error_msg"] = f"{type(e).__name__}: {e}"
    return result


# ============================ Excel 导出 ============================
def export_excel(results: list, out_path: str) -> None:
    """导出 Excel；目标文件被占用（如在 Excel 中打开）时自动另存为带时间戳的文件名。"""
    try:
        _export_excel_impl(results, out_path)
    except PermissionError:
        alt = out_path.replace(".xlsx", f"_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
        print(f"\n[警告] {out_path} 被占用（可能正在 Excel 中打开），结果已另存为：{alt}")
        _export_excel_impl(results, alt)


def _export_excel_impl(results: list, out_path: str) -> None:
    ok_results = [r for r in results if r["status"] == "ok"]
    failed = [r for r in results if r["status"] != "ok"]
    total = len(results)
    err_reports = [r for r in ok_results if r["has_error"]]
    err_rate = (len(err_reports) / len(ok_results)) if ok_results else 0.0

    # ---- 错误类型统计 ----
    type_counter = Counter(d["type"] for r in err_reports for d in r["error_details"])
    type_counts = sorted(type_counter.items(), key=lambda kv: (-kv[1], kv[0]))

    # ---- Sheet2 错误明细表数据（含无错报告，便于筛选核对；处理失败文件不计入）----
    rows = []
    for r in sorted(ok_results, key=lambda x: x["report_id"]):
        if r["has_error"]:
            details = r["error_details"] or [{
                "type": "、".join(r["error_categories"]) or "未分类",
                "original_text": "",
                "reason": "（模型未提供明细）",
                "suggestion": "",
            }]
            for d in details:
                rows.append([r["report_id"], "是", d["type"],
                             d["original_text"], d["reason"], d["suggestion"]])
        else:
            rows.append([r["report_id"], "否", "", "", "", ""])
    detail_df = pd.DataFrame(rows, columns=[
        "文件名/报告ID", "是否存在错误", "错误分类", "原文语句", "错误原因分析", "修改建议",
    ])

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        detail_df.to_excel(writer, sheet_name="错误明细表", index=False)
        wb = writer.book
        ws_over = wb.create_sheet("质控概览", 0)  # 插入为 Sheet1

        # ---- Sheet1 质控概览 ----
        title_font = Font(bold=True, size=14)
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="4472C4")

        ws_over["A1"] = "超声报告批量质控概览"
        ws_over["A1"].font = title_font

        ws_over["A3"], ws_over["B3"] = "指标", "数值"
        metrics = [
            ("总报告数", total),
            ("成功处理数", len(ok_results)),
            ("处理失败数", len(failed)),
            ("有错报告数", len(err_reports)),
            ("无错报告数", len(ok_results) - len(err_reports)),
        ]
        for i, (k, v) in enumerate(metrics):
            ws_over.cell(row=4 + i, column=1, value=k)
            ws_over.cell(row=4 + i, column=2, value=v)
        rate_row = 4 + len(metrics)
        ws_over.cell(row=rate_row, column=1, value="错误率")
        c = ws_over.cell(row=rate_row, column=2, value=err_rate)
        c.number_format = "0.00%"

        stat_header_row = rate_row + 2
        ws_over.cell(row=stat_header_row, column=1, value="错误类型")
        ws_over.cell(row=stat_header_row, column=2, value="错误数量（条）")

        if type_counts:
            for i, (cat, cnt) in enumerate(type_counts):
                ws_over.cell(row=stat_header_row + 1 + i, column=1, value=cat)
                ws_over.cell(row=stat_header_row + 1 + i, column=2, value=cnt)
            # 柱状图
            chart = BarChart()
            chart.type = "col"
            chart.title = "各错误类型数量统计"
            data = Reference(ws_over, min_col=2, min_row=stat_header_row,
                             max_row=stat_header_row + len(type_counts))
            cats_ref = Reference(ws_over, min_col=1, min_row=stat_header_row + 1,
                                 max_row=stat_header_row + len(type_counts))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats_ref)
            chart.legend = None
            chart.width, chart.height = 18, 9
            ws_over.add_chart(chart, "D3")
        else:
            ws_over.cell(row=stat_header_row + 1, column=1, value="（未发现错误）")

        for cell_ref in ("A3", "B3", f"A{stat_header_row}", f"B{stat_header_row}"):
            cell = ws_over[cell_ref]
            cell.font = head_font
            cell.fill = head_fill
        ws_over.column_dimensions["A"].width = 22
        ws_over.column_dimensions["B"].width = 18

        # ---- Sheet2 格式化 ----
        ws_det = writer.sheets["错误明细表"]
        for cell in ws_det[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        widths = [42, 12, 18, 55, 55, 55]
        for i, w in enumerate(widths, 1):
            ws_det.column_dimensions[get_column_letter(i)].width = w
        wrap = Alignment(wrap_text=True, vertical="top")
        for row in ws_det.iter_rows(min_row=2, max_row=max(ws_det.max_row, 2),
                                    min_col=4, max_col=6):
            for cell in row:
                cell.alignment = wrap
        ws_det.freeze_panes = "A2"
        if ws_det.max_row > 1:
            ws_det.auto_filter.ref = f"A1:F{ws_det.max_row}"


# ============================ 主流程 ============================
def main() -> None:
    parser = argparse.ArgumentParser(description="超声医学报告批量质控与错误分析")
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS,
                        help=f"并发线程数（默认 {DEFAULT_WORKERS}，建议 10-20）")
    parser.add_argument("--limit", type=int, default=0,
                        help="仅处理前 N 份报告（0=全部），可用于小样本试跑控制成本")
    parser.add_argument("--model", default=MODEL_NAME, help="DeepSeek 模型名称")
    parser.add_argument("--out", default=OUTPUT_XLSX, help="输出 Excel 路径")
    args = parser.parse_args()

    if not API_KEY or API_KEY.startswith("sk-请"):
        print("错误：未配置 API_KEY，请设置环境变量 DEEPSEEK_API_KEY 或修改脚本配置区。")
        sys.exit(1)

    all_files = sorted(REPORT_DIR.glob("*.txt"))
    if not all_files:
        print(f"错误：在 {REPORT_DIR.resolve()} 下未找到任何 .txt 报告文件。")
        sys.exit(1)
    files = all_files
    if args.limit > 0:
        files = files[:args.limit]

    print(f"共发现 {len(all_files)} 份报告，本次处理 {len(files)} 份 | "
          f"模型: {args.model} | 并发: {args.workers}")

    # max_retries=0：关闭 SDK 内置重试，统一由 call_with_retry 控制
    client = OpenAI(api_key=API_KEY, base_url=BASE_URL,
                    timeout=REQUEST_TIMEOUT, max_retries=0)

    results = []
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = {pool.submit(qc_one_report, client, f, args.model): f for f in files}
        with tqdm(total=len(futures), desc="质控审核进度", unit="份", ncols=80) as pbar:
            for fut in as_completed(futures):
                r = fut.result()
                results.append(r)
                if r["status"] != "ok":
                    tqdm.write(f"[失败] {r['report_id']}: {r['error_msg']}")
                pbar.update(1)

    export_excel(results, args.out)

    # ---- 控制台汇总 ----
    ok_results = [r for r in results if r["status"] == "ok"]
    failed = [r for r in results if r["status"] != "ok"]
    err_reports = [r for r in ok_results if r["has_error"]]
    print("\n========== 质控完成 ==========")
    print(f"总报告数: {len(results)} | 成功: {len(ok_results)} | 失败: {len(failed)}")
    print(f"有错报告: {len(err_reports)} | 错误率: "
          f"{len(err_reports) / len(ok_results):.2%}" if ok_results else "无成功处理记录")
    print(f"结果已导出: {Path(args.out).resolve()}")
    if failed:
        print(f"\n处理失败的 {len(failed)} 份文件（可重新运行仅处理这些文件）：")
        for r in failed[:20]:
            print(f"  - {r['report_id']}: {r['error_msg'][:120]}")
        if len(failed) > 20:
            print(f"  ... 等共 {len(failed)} 份")


if __name__ == "__main__":
    main()
